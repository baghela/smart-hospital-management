"""Document text extraction and semantic-boundary chunking.

Supports PDF (PyMuPDF with OCR fallback hook), DOCX, CSV, TXT and Markdown.
Page numbers and detected section headings are preserved for citations.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import List, Tuple

from backend.core.config import get_settings
from backend.core.errors import ValidationError

_HEADING_RE = re.compile(r"^\s*(?:#{1,4}\s+|\d+(?:\.\d+)*\s+)?([A-Z][A-Za-z0-9 ,&/()\-]{3,80})\s*:?\s*$")


@dataclass
class PageText:
    page: int
    text: str


@dataclass
class Chunk:
    index: int
    section: str
    page: int
    content: str

    @property
    def token_estimate(self) -> int:
        return max(1, len(self.content) // 4)


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------
def extract_pages(filename: str, data: bytes) -> List[PageText]:
    ext = "." + filename.rsplit(".", 1)[-1].lower()
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext == ".csv":
        return _extract_csv(data)
    if ext in (".txt", ".md"):
        return _paginate(data.decode("utf-8", errors="replace"))
    raise ValidationError(f"Unsupported document type '{ext}'.")


def _paginate(text: str, chars_per_page: int = 3000) -> List[PageText]:
    pages = [text[i : i + chars_per_page] for i in range(0, max(len(text), 1), chars_per_page)]
    return [PageText(page=i + 1, text=p) for i, p in enumerate(pages) if p.strip()]


def _extract_pdf(data: bytes) -> List[PageText]:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("PyMuPDF is required to ingest PDF files.") from exc

    pages: List[PageText] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for number, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if not text:
                text = _ocr_page(page)
            if text:
                pages.append(PageText(page=number, text=text))
    if not pages:
        raise ValidationError("No extractable text found in the PDF (scanned file without OCR).")
    return pages


def _ocr_page(page) -> str:
    """OCR fallback for scanned PDFs; returns '' when OCR is unavailable."""
    try:  # pragma: no cover - depends on local tesseract install
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(image).strip()
    except Exception:
        return ""


def _extract_docx(data: bytes) -> List[PageText]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("python-docx is required to ingest DOCX files.") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _paginate("\n".join(parts))


def _extract_csv(data: bytes) -> List[PageText]:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise ValidationError("CSV file is empty.")
    header = rows[0]
    lines = []
    for row in rows[1:]:
        pairs = [f"{h.strip()}: {v.strip()}" for h, v in zip(header, row) if v.strip()]
        if pairs:
            lines.append(" | ".join(pairs))
    return _paginate("\n".join(lines))


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------
def detect_section(line: str, current: str) -> str:
    match = _HEADING_RE.match(line)
    if match and len(line.split()) <= 12:
        return match.group(1).strip()
    return current


def chunk_pages(
    pages: List[PageText], *, chunk_chars: int | None = None, overlap: int | None = None
) -> List[Chunk]:
    """Chunk on paragraph/section boundaries, keeping page + section metadata."""
    settings = get_settings()
    size = chunk_chars or settings.chunk_chars
    slide = overlap if overlap is not None else settings.chunk_overlap

    chunks: List[Chunk] = []
    section = "General"
    index = 0

    for page in pages:
        buffer: List[str] = []
        buffer_len = 0
        section_at_start = section
        for raw_line in page.text.splitlines():
            line = raw_line.rstrip()
            new_section = detect_section(line, section)
            boundary = new_section != section and buffer
            if boundary or buffer_len + len(line) > size:
                content = "\n".join(buffer).strip()
                if content:
                    chunks.append(Chunk(index, section_at_start, page.page, content))
                    index += 1
                tail = content[-slide:] if slide and content else ""
                buffer = [tail] if tail else []
                buffer_len = len(tail)
                section_at_start = new_section
            section = new_section
            if line.strip():
                buffer.append(line)
                buffer_len += len(line) + 1
        content = "\n".join(buffer).strip()
        if content:
            chunks.append(Chunk(index, section_at_start, page.page, content))
            index += 1
    return [c for c in chunks if len(c.content) > 40]


def extract_and_chunk(filename: str, data: bytes) -> Tuple[List[PageText], List[Chunk]]:
    pages = extract_pages(filename, data)
    return pages, chunk_pages(pages)

